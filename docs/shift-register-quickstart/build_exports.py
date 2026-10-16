"""Build the short shift-register lesson as editable Word and PowerPoint files.

Usage: python build_exports.py [--export-pdf] [--qa]
Dependencies: python-docx, python-pptx, Pillow. Optional PDF review: PyMuPDF.
PDF export uses installed Microsoft Word and PowerPoint on Windows.
The lesson.json file is the content source; this builder contains no old course text.
"""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import textwrap

BASE = Path(__file__).resolve().parent
REPO = BASE.parents[1]
WORK = REPO.parent / '.codex-work' / 'ee6019'
for candidate in (WORK / 'slides' / 'deps', WORK / 'beginner-guide' / 'deps'):
    if candidate.is_dir():
        sys.path.insert(0, str(candidate))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR

NAVY = '16324F'
TEAL = '087F8C'
ORANGE = 'D96A2B'
INK = '24384B'
MUTED = '5E6D78'
PAPER = 'FBF8F1'
PALE = 'EAF4F3'
LINE = 'D8E0E1'
WHITE = 'FFFFFF'
FONT = 'Calibri'
REVIEW = WORK / 'shift-register-review'
OUTPUT_SUFFIX = ''

def output_file(name):
    p=Path(name)
    return BASE/(p.stem+OUTPUT_SUFFIX+p.suffix)

# Slide wording is deliberately shorter than the self-study guide. Full instructions
# are retained in the matching guide page and in editable presenter notes.
SLIDE_SUMMARIES = {
    'start': ('Bring your working Verilog design. Leave with a tested Tiny Tapeout layout.', [
        'Open the HTML guide and the starter ZIP beside GitHub.',
        'Follow guide pages 2–14, including GLS. Tick each checkpoint.']),
    'design': ('Eight stored bits. One input bit. One clock edge at a time.', [
        'Extract the starter ZIP. Open src/shift_register.v.',
        'Reset clears; enable shifts; otherwise hold.',
        'Shift 1, 0, 1 in the HTML clock experiment.']),
    'repo': ('Make your own copy of the official SKY Verilog template.', [
        'Open TinyTapeout/ttsky-verilog-template.',
        'Use this template → Create a new repository.',
        'Choose your account. Name: tt-shift-register. Visibility: Public.']),
    'settings': ('Actions runs your tools. Pages shows the layout.', [
        'Open Actions. Enable workflows if asked.',
        'Settings → Pages → Source: GitHub Actions.',
        'Keep .github, src/config.json and test/requirements.txt.']),
    'pins': ('A wrapper connects your signal names to Tiny Tapeout’s fixed ports.', [
        'ui_in[0] → serial input. ui_in[1] → shift enable.',
        'Clock connects directly. Invert rst_n to make the core reset.',
        'q → uo_out. Set unused uio_out and uio_oe to zero.']),
    'wrapper': ('Keep the existing circuit. Add one small module around it.', [
        'Replace src/project.v with the supplied wrapper.',
        'Keep tt_um_student_shift_register for this exercise.',
        'The instance below joins each circuit port to its chip signal.']),
    'files': ('Both tools need both files. Three top-module names must agree.', [
        'One top name: tt_um_student_shift_register.',
        'Two source lists: project.v and shift_register.v.',
        'Check info.yaml, test/Makefile and test/tb.v.']),
    'metadata': ('Describe how to use the circuit. Keep its clock settings consistent.', [
        'info.yaml: your author name, title and correct pin labels.',
        'docs/info.md: behaviour, test procedure and hardware needs.',
        '50 MHz target = 20 ns per clock cycle.']),
    'test': ('Test your shift register through its external pins.', [
        'Replace test/test.py with the supplied test.',
        'Check reset, shift, hold, old data leaving and reset priority.',
        'Compare each observed output with the expected value.']),
    'upload': ('Upload the contents of the extracted starter into your repository root.', [
        'Code → Add file → Upload files.',
        'Drag info.yaml plus src, test and docs. Check all seven paths.',
        'Commit to main: Add shift register. If needed, create and merge a pull request.']),
    'rtl': ('Read the test result for the commit containing all seven files.', [
        'Actions → test → your final commit’s run.',
        'Open the test job → Run tests. Zero failures or errors.',
        'If it fails, fix the first error and save a new commit.']),
    'gds': ('Verilog becomes physical gates and wires. Check every required job.', [
        'docs: download PDF/datasheet and check the pin table.',
        'gds: check build, precheck, gate-level test and viewer.',
        'Open the layout. Download tt_submission and GDS_logs.']),
    'gls-prepare': ('GLS runs your test against generated logic cells and their connections.', [
        'Keep the same test/test.py and external port checks.',
        'Keep the GL_TEST fixture’s VPWR and VGND connections.',
        'The cloud job supplies the netlist and cell models; GATES=yes selects them.']),
    'gls-run': ('Open the gate-level job for the same commit as your RTL test.', [
        'Actions → gds → your run → gl_test → GL test → Run tests.',
        'Read TESTS=1 PASS=1 FAIL=0 and confirm the whole job succeeds.',
        'Summary → Artifacts → gatelevel_test_results: save results.xml and waveform.']),
    'finish': ('Save one identifiable result: circuit, checks, layout and evidence.', [
        'Keep repository, commit and workflow links together.',
        'Save the GDS, reports and datasheet for that same commit.',
        'For your next design: map → wrap → list sources → test → build.']),
}


def rgb(value):
    return tuple(int(value[i:i+2], 16) for i in (0, 2, 4))


def font(size, bold=False):
    path = Path('C:/Windows/Fonts/calibrib.ttf' if bold else 'C:/Windows/Fonts/calibri.ttf')
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def diagram_items(kind):
    if kind == 'route':
        return [('Your Verilog', 'existing design'), ('Connect pins', 'Tiny Tapeout wrapper'),
                ('Test', 'check the behaviour'), ('Build layout', 'GDS + checks'),
                ('Save evidence', 'ready for review')]
    if kind == 'checks':
        return [('RTL test', 'behaviour'), ('GDS build', 'physical layout'),
                ('Precheck', 'template rules'), ('Gate-level test', 'built circuit'),
                ('Archive', 'same final commit')]
    if kind == 'files':
        return [('src/', 'Verilog design + wrapper'), ('test/', 'test + build list'),
                ('info.yaml', 'project settings'), ('docs/info.md', 'how to use the chip')]
    return []


def draw_diagram(kind, dest):
    """High resolution guide diagram. Slides use separate editable shapes."""
    w, h = 1800, (740 if kind=='checks' else 420)
    im = Image.new('RGB', (w, h), '#' + PAPER)
    d = ImageDraw.Draw(im)
    def centred(box, lines):
        x0, y0, x1, y1 = box
        for i, (line, size, colour, bold) in enumerate(lines):
            ft = font(size, bold)
            tw = d.textbbox((0, 0), line, font=ft)[2]
            d.text(((x0+x1-tw)/2, y0+42+i*58), line, font=ft, fill='#'+colour)
    def arrow(x0, y, x1):
        d.line((x0, y, x1-13, y), fill='#'+TEAL, width=7)
        d.polygon([(x1, y), (x1-20, y-13), (x1-20, y+13)], fill='#'+TEAL)
    if kind == 'route':
        items = diagram_items(kind)
        count = len(items)
        gap = 45
        boxw = (w-70-(count-1)*gap)/count
        for i, (title, sub) in enumerate(items):
            x = 35+i*(boxw+gap)
            d.rounded_rectangle((x, 105, x+boxw, 285), radius=22, fill='#'+WHITE, outline='#'+LINE, width=3)
            centred((x, 105, x+boxw, 285), [(title, 35, NAVY, True), (sub, 25, MUTED, False)])
            if i < count-1:
                arrow(x+boxw+7, 196, x+boxw+gap-7)
        d.text((45,335),'Recommended learning order • test, docs and gds workflows start independently.',font=font(28),fill='#'+MUTED)
    elif kind == 'files':
        items=[('src/project.v','module name'),('info.yaml','top_module'),('test/tb.v','instance name')]
        for i,(title,sub) in enumerate(items):
            x=35+i*600
            d.rounded_rectangle((x,35,x+555,220),radius=22,fill='#'+WHITE,outline='#'+LINE,width=3)
            centred((x,35,x+555,220),[(title,35,NAVY,True),(sub,28,MUTED,False)])
            if i<2: arrow(x+562,126,x+594)
        d.text((45,253),'ONE TOP NAME:  tt_um_student_shift_register',font=font(34,True),fill='#'+TEAL)
        d.text((45,320),'TWO SOURCE LISTS:  info.yaml + test/Makefile → project.v + shift_register.v',font=font(31),fill='#'+INK)
    elif kind == 'checks':
        nodes=[(15,265,350,125,'Commit / push','same saved version'),
               (500,20,450,125,'RTL test','source behaviour'),
               (500,220,450,125,'docs','build the datasheet'),
               (500,420,450,125,'GDS build','place gates + route wires'),
               (1220,250,550,115,'Precheck','integration rules'),
               (1220,410,550,115,'Gate-level test (GLS)','same checks on logic cells'),
               (1220,570,550,115,'Viewer','publish the layout view')]
        for xx,yy,bw,bh,title,sub in nodes:
            d.rounded_rectangle((xx,yy,xx+bw,yy+bh),radius=18,fill='#'+WHITE,outline='#'+LINE,width=3)
            ft=font(32,True); tw=d.textbbox((0,0),title,font=ft)[2];d.text((xx+(bw-tw)/2,yy+22),title,font=ft,fill='#'+NAVY)
            ft=font(27);tw=d.textbbox((0,0),sub,font=ft)[2];d.text((xx+(bw-tw)/2,yy+69),sub,font=ft,fill='#'+MUTED)
        d.line((365,327,425,327),fill='#'+TEAL,width=5);d.line((425,82,425,482),fill='#'+TEAL,width=5)
        for yy in [82,282,482]:arrow(425,yy,490)
        d.line((950,482,1080,482),fill='#'+TEAL,width=5);d.line((1080,307,1080,627),fill='#'+TEAL,width=5)
        for yy in [307,467,627]:arrow(1080,yy,1210)
        d.text((35,705),'Three independent workflows; precheck, GLS and viewer wait for the GDS build.',font=font(27),fill='#'+MUTED)
    elif kind == 'shift':
        d.text((40, 70), 'serial input', font=font(34, True), fill='#'+NAVY)
        d.text((40, 120), 'one new bit', font=font(29), fill='#'+MUTED)
        arrow(228, 198, 303)
        for i in range(8):
            x = 315+i*160
            d.rounded_rectangle((x, 134, x+135, 263), radius=16, fill='#'+WHITE, outline='#'+TEAL, width=4)
            d.text((x+39, 166), f'q[{i}]', font=font(36, True), fill='#'+NAVY)
            if i < 7:
                arrow(x+139, 199, x+157)
        d.text((315, 309), 'At each rising clock edge, enabled bits move one place to the right.', font=font(32), fill='#'+INK)
    elif kind == 'mapping':
        left = [('ui_in[0]', 'serial input'), ('ui_in[1]', 'shift enable'), ('clk', 'clock'), ('rst_n', 'reset (active low)')]
        for i,(pin,meaning) in enumerate(left):
            yy = 34+i*84
            d.rounded_rectangle((30, yy, 477, yy+68), radius=12, fill='#'+WHITE, outline='#'+LINE, width=2)
            d.text((52, yy+13), pin, font=font(32, True), fill='#'+NAVY)
            d.text((220, yy+16), meaning, font=font(27), fill='#'+MUTED)
            arrow(490, yy+34, 642)
            if pin=='rst_n':
                d.ellipse((548,yy+12,605,yy+58),fill='#'+PAPER,outline='#'+ORANGE,width=3)
                d.text((562,yy+3),'~',font=font(43,True),fill='#'+ORANGE)
        d.rounded_rectangle((655, 44, 1185, 365), radius=25, fill='#'+PALE, outline='#'+TEAL, width=3)
        d.text((738, 139), 'your shift register', font=font(42, True), fill='#'+NAVY)
        d.text((727, 202), 'rst = ~rst_n', font=font(30), fill='#'+MUTED)
        arrow(1202, 202, 1310)
        d.rounded_rectangle((1325, 129, 1770, 273), radius=20, fill='#'+WHITE, outline='#'+LINE, width=3)
        centred((1325, 129, 1770, 273), [('uo_out[7:0]', 38, NAVY, True), ('eight stored bits', 28, MUTED, False)])
    im.save(dest)


def xml_shade(cell, colour):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), colour)
    props.append(shade)


def prevent_row_split(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement('w:cantSplit')
    props.append(element)


def set_cell_margins(cell, top=90, start=140, bottom=90, end=140):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for tag, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el = OxmlElement('w:'+tag)
        el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)


def para(document, text='', size=11, bold=False, colour=INK, before=0, after=6, keep=False):
    p = document.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before); fmt.space_after = Pt(after)
    fmt.line_spacing = 1.07
    fmt.keep_with_next = keep
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(colour)
    return p


def hyperlink(p, label, url, size=9.5):
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'), TEAL); props.append(col)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size*2))); props.append(sz)
    un = OxmlElement('w:u'); un.set(qn('w:val'), 'single'); props.append(un)
    run.append(props); tx = OxmlElement('w:t'); tx.text=label; run.append(tx); link.append(run); p._p.append(link)


def box(document, label, text, colour=TEAL, size=10.5):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0,0); xml_shade(cell, PALE if colour==TEAL else 'FFF0E3')
    set_cell_margins(cell, 125, 170, 125, 170)
    p = cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
    r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(colour); r.font.size=Pt(size)
    r=p.add_run(text); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(INK)
    prevent_row_split(table.rows[0])
    para(document, after=1, size=2)


def code_box(document, code):
    para(document, code.get('file','Verilog'), size=9, bold=True, colour=TEAL, before=3, after=4, keep=True)
    table=document.add_table(rows=1,cols=1); cell=table.cell(0,0)
    xml_shade(cell, 'F0F3F5'); set_cell_margins(cell, 90, 150, 90, 150)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    lines=code['text'].rstrip().splitlines()
    code_size=9 if len(lines)<=25 else 8.5
    p.paragraph_format.line_spacing=Pt(code_size*1.12)
    r=p.add_run('\n'.join(lines)); r.font.name='Consolas'; r.font.size=Pt(code_size); r.font.color.rgb=RGBColor.from_string(INK)
    prevent_row_split(table.rows[0]); para(document,size=2,after=1)


def image_path(src):
    p=Path(src)
    return p if p.is_absolute() else BASE/p


def display_picture(src):
    """Enlarge a genuine UI detail; retain the unaltered source in images/."""
    source=image_path(src)
    if source.name!='pages.png': return source
    dest=REVIEW/'pages-source-detail.png'
    with Image.open(source) as im:
        im.crop((int(im.width*.25),int(im.height*.59),int(im.width*.69),int(im.height*.86))).save(dest)
    return dest


def write_doc(lesson, assets):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.55)
    sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)
    sec.header_distance=Inches(.23); sec.footer_distance=Inches(.22)
    normal=doc.styles['Normal']; normal.font.name=FONT; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5)
    h=sec.header.paragraphs[0]; h.text='TINY TAPEOUT  /  SHIFT REGISTER QUICKSTART'
    h.runs[0].font.name=FONT; h.runs[0].font.size=Pt(8); h.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=foot.add_run('EE6019  •  '+lesson.get('date','')+'  •  '); run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(MUTED)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
    for i,page in enumerate(lesson['pages']):
        if i: doc.add_page_break()
        para(doc,page.get('kicker',f'STEP {i}'),size=9,bold=True,colour=TEAL,after=6)
        p=para(doc,page['title'],size=25 if len(page['title'])<50 else 22,bold=True,colour=NAVY,after=10,keep=True)
        p.style=doc.styles['Heading 1']
        for r in p.runs: r.font.name=FONT; r.font.size=Pt(25 if len(page['title'])<50 else 22); r.font.color.rgb=RGBColor.from_string(NAVY)
        if page.get('intro'): para(doc,page['intro'],size=11,after=9)
        if page.get('diagram'):
            doc.add_picture(str(assets[page['diagram']]),width=Inches(6.8))
            doc.paragraphs[-1].paragraph_format.space_after=Pt(5)
        for n,action in enumerate(page.get('actions',[]),1):
            p=para(doc,size=10.8,after=7)
            r=p.add_run(f'{n:02d}   '); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
            r=p.add_run(action); r.font.size=Pt(10.8); r.font.color.rgb=RGBColor.from_string(INK)
        if page.get('table'):
            data=page['table']; tab=doc.add_table(rows=1,cols=len(data['headers']))
            tab.alignment=WD_TABLE_ALIGNMENT.CENTER
            for j,head in enumerate(data['headers']):
                cell=tab.rows[0].cells[j]; xml_shade(cell, NAVY); set_cell_margins(cell)
                r=cell.paragraphs[0].add_run(head); r.font.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor.from_string(WHITE)
            for n,row in enumerate(data['rows']):
                cells=tab.add_row().cells
                for j,txt in enumerate(row):
                    xml_shade(cells[j], PALE if n%2==0 else 'F6F7F6'); set_cell_margins(cells[j])
                    p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
                    r=p.add_run(str(txt)); r.font.size=Pt(9.5)
                prevent_row_split(tab.rows[-1])
            para(doc,size=3,after=2)
        if page.get('code'): code_box(doc,page['code'])
        if page.get('image'):
            data=page['image']; pic=display_picture(data['src'])
            with Image.open(pic) as im: ratio=im.height/im.width
            maxh=2.65 if page.get('code') or page.get('table') else 3.0
            width=min(6.8,maxh/ratio)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
            p.add_run().add_picture(str(pic),width=Inches(width))
            if data.get('caption'):
                caption=('Enlarged detail. '+data['caption']) if Path(data['src']).name=='pages.png' else data['caption']
                para(doc,caption,size=8.5,colour=MUTED,after=7)
        for term in page.get('terms',[]):
            p=para(doc,size=9.6,after=4)
            r=p.add_run(term['word']+' = '); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
            r=p.add_run(term['meaning']); r.font.size=Pt(9.6)
        if page.get('predict'):
            pred=page['predict']; box(doc,'PAUSE & PREDICT',pred['question'],ORANGE)
            para(doc,'Answer: '+pred['answer'],size=9.5,colour=MUTED,after=7)
        for help in page.get('help',[]):
            para(doc,help['problem']+' → '+help['fix'],size=9.5,after=5)
        if page.get('checkpoint'): box(doc,'CHECK BEFORE YOU CONTINUE',page['checkpoint'])
        links=page.get('links',[])
        if links:
            p=para(doc,size=9,after=0)
            for j,link in enumerate(links):
                if j: p.add_run('   •   ')
                hyperlink(p,link['label'],link['url'])
    doc.core_properties.title=lesson['title']
    doc.core_properties.subject='Short, practical route from an existing Verilog design to a checked Tiny Tapeout layout'
    doc.core_properties.author='EE6019 Research Project'
    doc.save(output_file('quickstart.docx'))


def ppt_text(slide,x,y,w,h,text,size=20,colour=INK,bold=False,name=None):
    shape=slide.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h))
    if name: shape.name=name
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for n,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
        p.text=line; p.font.name=FONT; p.font.size=PP(size); p.font.bold=bold; p.font.color.rgb=PC.from_string(colour)
        p.space_after=PP(8)
    return shape


def ppt_rect(slide,x,y,w,h,fill=WHITE,line=LINE,radius=True):
    typ=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh=slide.shapes.add_shape(typ,PI(x),PI(y),PI(w),PI(h)); sh.fill.solid(); sh.fill.fore_color.rgb=PC.from_string(fill)
    sh.line.color.rgb=PC.from_string(line); sh.line.width=PP(1)
    if radius: sh.adjustments[0]=.08
    return sh


def ppt_arrow(slide,x,y,w):
    sh=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,PI(x),PI(y-.07),PI(w),PI(.14))
    sh.fill.solid(); sh.fill.fore_color.rgb=PC.from_string(TEAL); sh.line.fill.background()


def ppt_diagram(slide,kind,x,y,w,h):
    if kind=='route':
        items=diagram_items(kind); gap=.2; bw=(w-gap*(len(items)-1))/len(items)
        for n,(title,sub) in enumerate(items):
            xx=x+n*(bw+gap)
            ppt_rect(slide,xx,y,bw,h,WHITE)
            ppt_text(slide,xx+.12,y+.18,bw-.24,.6,title,19,NAVY,True)
            ppt_text(slide,xx+.12,y+.84,bw-.24,h-.9,sub,16,MUTED)
            if n<len(items)-1: ppt_arrow(slide,xx+bw+.02,y+h/2,gap-.04)
        ppt_text(slide,x,y+h+.035,w,.24,'Recommended learning order • test, docs and gds workflows start independently.',11,MUTED)
    elif kind=='files':
        items=[('src/project.v','module name'),('info.yaml','top_module'),('test/tb.v','instance name')]
        gap=.24; bw=(w-2*gap)/3
        for n,(title,sub) in enumerate(items):
            xx=x+n*(bw+gap); ppt_rect(slide,xx,y,bw,h*.7,WHITE)
            ppt_text(slide,xx+.1,y+.12,bw-.2,.38,title,18,NAVY,True)
            ppt_text(slide,xx+.1,y+.52,bw-.2,.35,sub,16,MUTED)
            if n<2: ppt_arrow(slide,xx+bw+.03,y+h*.36,gap-.06)
        ppt_text(slide,x,y+h*.77,w,.45,'Both source lists: project.v + shift_register.v',18,TEAL,True)
    elif kind=='checks':
        # Common trigger branches into independent workflows. GDS fans out again.
        def node(xx,yy,ww,title,sub):
            ppt_rect(slide,xx,yy,ww,.62,WHITE)
            ppt_text(slide,xx+.08,yy+.06,ww-.16,.29,title,17,NAVY,True)
            ppt_text(slide,xx+.08,yy+.34,ww-.16,.23,sub,12,MUTED)
        startx=x; midx=x+w*.28; endx=x+w*.67
        node(startx,y+1.0,w*.20,'Commit / push','same saved version')
        for yy,title,sub in [(y,'RTL test','source behaviour'),(y+.85,'docs','build the datasheet'),(y+1.70,'GDS build','gates + wires')]:node(midx,yy,w*.25,title,sub)
        for yy,title,sub in [(y+.65,'Precheck','integration rules'),(y+1.5,'Gate-level test (GLS)','same test on logic cells'),(y+2.35,'Viewer','publish layout')]:node(endx,yy,w*.32,title,sub)
        def line(x1,y1,x2,y2):
            c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PI(x1),PI(y1),PI(x2),PI(y2));c.line.color.rgb=PC.from_string(TEAL);c.line.width=PP(1.5)
        spine=x+w*.235
        line(startx+w*.20,y+1.31,spine,y+1.31);line(spine,y+.31,spine,y+2.01)
        for yy in [y+.31,y+1.16,y+2.01]:ppt_arrow(slide,spine,yy,midx-spine-.03)
        spine2=x+w*.60
        line(midx+w*.25,y+2.01,spine2,y+2.01);line(spine2,y+.96,spine2,y+2.66)
        for yy in [y+.96,y+1.81,y+2.66]:ppt_arrow(slide,spine2,yy,endx-spine2-.03)
    elif kind=='shift':
        ppt_text(slide,x,y+.06,w,.4,'serial input → eight stored bits',19,TEAL,True)
        gap=.12; bw=(w-.65-gap*7)/8; xx=x+.65
        ppt_text(slide,x,y+.75,.58,.6,'in',20,NAVY,True)
        ppt_arrow(slide,x+.29,y+1.35,.3)
        for n in range(8):
            px=xx+n*(bw+gap); ppt_rect(slide,px,y+.72,bw,.95,WHITE,TEAL)
            ppt_text(slide,px+.04,y+1.03,bw-.08,.35,f'q{n}',18,NAVY,True)
            if n<7: ppt_arrow(slide,px+bw+.01,y+1.2,gap-.02)
        ppt_text(slide,x,y+1.9,w,.9,'Enabled clock edge: each bit moves one position.',19,MUTED)
    elif kind=='mapping':
        pins=[('ui_in[0]','serial input'),('ui_in[1]','enable'),('clk','clock'),('rst_n','active-low reset')]
        lw=w*.35; cw=w*.30; rw=w*.25
        for n,(p,sub) in enumerate(pins):
            yy=y+n*.64; ppt_rect(slide,x,yy,lw,.51,WHITE)
            ppt_text(slide,x+.09,yy+.1,lw-.18,.35,p+'  '+sub,15,NAVY)
            ppt_arrow(slide,x+lw+.04,yy+.25,w*.055)
            if p=='rst_n':
                ppt_text(slide,x+lw+.02,yy+.05,w*.07,.47,'~',24,ORANGE,True)
        cx=x+w*.42; ppt_rect(slide,cx,y+.25,cw,1.9,PALE,TEAL)
        ppt_text(slide,cx+.13,y+.76,cw-.26,1.0,'Your shift\nregister',21,NAVY,True)
        ppt_arrow(slide,cx+cw+.04,y+1.2,w*.04)
        rx=x+w*.79; ppt_rect(slide,rx,y+.71,rw,1.12,WHITE)
        ppt_text(slide,rx+.08,y+.91,rw-.16,.8,'uo_out\n8 bits',18,NAVY,True)


def new_slide(prs,page,index,detail=False):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=PC.from_string(PAPER)
    ppt_rect(slide,0,0,.16,7.5,TEAL,TEAL,False)
    ppt_text(slide,.56,.29,11.8,.32,page.get('kicker','TINY TAPEOUT QUICKSTART')+('  /  DETAIL' if detail else ''),12,TEAL,True)
    title=page['title']+(' — in practice' if detail and len(page['title'])<42 else '')
    ppt_text(slide,.56,.84,12.0,.96,title,31 if len(title)<62 else 27,NAVY,True)
    ppt_text(slide,.58,7.05,10.9,.2,'SHIFT REGISTER QUICKSTART  •  Guide page '+str(index+1),10,MUTED)
    if page.get('links'):
        link=page['links'][0]
        button=ppt_text(slide,8.7,7.005,3.04,.30,'Open reference ↗',12,TEAL,True)
        button.text_frame.paragraphs[0].runs[0].hyperlink.address=link['url']
    ppt_text(slide,12.1,7.0,.64,.3,str(len(prs.slides)),11,MUTED)
    return slide


def ppt_actions(slide,actions,x,y,w,size=20,maxh=3.8):
    # Each action gets its own native textbox and visible step number.
    if not actions: return
    weights=[max(1,len(a)/(w*5.7)) for a in actions]
    unit=maxh/sum(weights)
    for n,(action,weight) in enumerate(zip(actions,weights),1):
        hh=unit*weight
        ppt_text(slide,x,y,.42,.4,f'{n:02d}',size,TEAL,True)
        ppt_text(slide,x+.52,y,w-.50,hh-.12,action,size,INK)
        y+=hh


def ppt_checkpoint(slide,text):
    ppt_rect(slide,.55,6.2,12.2,.66,PALE,PALE)
    ppt_text(slide,.72,6.31,1.14,.33,'CHECK',14,TEAL,True)
    ppt_text(slide,1.92,6.30,10.57,.43,text,16,INK)


def code_excerpt(code,max_lines=8):
    lines=code['text'].rstrip().splitlines()
    if len(lines)<=max_lines: return '\n'.join(lines),False
    # Show the behavioural rule or instance connections, with no long commentary.
    if any('always @' in l for l in lines):
        start=next(i for i,l in enumerate(lines) if 'always @' in l)
        selected=lines[start:start+max_lines]
    else:
        start=next((i for i,l in enumerate(lines) if 'shift_register' in l and '(' in l and 'module' not in l),0)
        selected=lines[start:start+max_lines]
    selected=[l.split('//')[0].rstrip() for l in selected]
    while selected and not selected[-1]: selected.pop()
    return '\n'.join(selected),True


def ppt_code(slide,code,x,y,w,h):
    ppt_rect(slide,x,y,w,h,'F0F3F5','E0E6E8')
    ppt_text(slide,x+.17,y+.13,w-.34,.35,code.get('file','Code'),14,TEAL,True)
    text,shortened=code_excerpt(code)
    shape=ppt_text(slide,x+.17,y+.60,w-.34,h-.93,text,14,INK)
    for p in shape.text_frame.paragraphs:
        p.font.name='Consolas'; p.space_after=PP(0)
    shape.text_frame.word_wrap=False
    if shortened: ppt_text(slide,x+.17,y+h-.32,w-.34,.26,'Excerpt • full file in the guide and example download',11,MUTED)


def ppt_picture(slide,data,x,y,w,h):
    src=display_picture(data['src'])
    with Image.open(src) as im: ratio=im.height/im.width
    ww=min(w,(h-.5)/ratio); hh=ww*ratio
    slide.shapes.add_picture(str(src),PI(x+(w-ww)/2),PI(y),width=PI(ww),height=PI(hh))
    if data.get('caption'):
        caption=('Enlarged detail. '+data['caption']) if Path(data['src']).name=='pages.png' else data['caption']
        ppt_text(slide,x,y+hh+.12,w,.54,caption,12,MUTED)


def ppt_table(slide,data,x,y,w,h):
    nrows=len(data['rows'])+1; ncols=len(data['headers'])
    table=slide.shapes.add_table(nrows,ncols,PI(x),PI(y),PI(w),PI(h)).table
    for r,row in enumerate([data['headers']]+data['rows']):
        for c,txt in enumerate(row):
            cell=table.cell(r,c); cell.text=str(txt); cell.margin_left=PI(.12); cell.margin_right=PI(.08)
            cell.margin_top=PI(.07); cell.margin_bottom=PI(.05)
            cell.fill.solid(); cell.fill.fore_color.rgb=PC.from_string(NAVY if r==0 else PALE if r%2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT; p.font.size=PP(16 if nrows>6 else 18); p.font.bold=r==0
                p.font.color.rgb=PC.from_string(WHITE if r==0 else INK); p.space_after=PP(0)


def notes(slide,page,guidepage):
    content=[f'Guide page {guidepage}: {page["title"]}',page.get('intro','')]
    content += page.get('actions',[])
    content += [t['word']+' = '+t['meaning'] for t in page.get('terms',[])]
    content += [h['problem']+' → '+h['fix'] for h in page.get('help',[])]
    if page.get('predict'): content += ['Pause and predict: '+page['predict']['question'],'Answer: '+page['predict']['answer']]
    if page.get('checkpoint'): content.append('Check: '+page['checkpoint'])
    content += [l['label']+': '+l['url'] for l in page.get('links',[])]
    if page.get('code'): content += ['Full code ('+page['code'].get('file','')+'):',page['code']['text']]
    slide.notes_slide.notes_text_frame.text='\n\n'.join(c for c in content if c)


def write_slides(lesson):
    prs=Presentation(); prs.slide_width=PI(13.333); prs.slide_height=PI(7.5)
    for i,page in enumerate(lesson['pages']):
        slide=new_slide(prs,page,i)
        notes(slide,page,i+1)
        intro,actions=SLIDE_SUMMARIES.get(page.get('id'),(page.get('intro',''),page.get('actions',[])))
        hasvisual=page.get('code') or page.get('image') or page.get('table') or page.get('diagram')
        if intro: ppt_text(slide,.58,1.89,12.03,.78,intro,19,MUTED)
        y=2.85
        if page.get('diagram')=='checks':
            ppt_diagram(slide,'checks',.58,2.70,12.1,3.0)
            ppt_text(slide,.65,5.85,12.0,.42,'Check every required job for one commit. The GLS walkthrough follows.',19,TEAL,True)
        elif page.get('diagram') and page['diagram'] in ('route','files'):
            ppt_diagram(slide,page['diagram'],.58,y,12.1,1.43)
            if actions: ppt_actions(slide,actions[:2],.62,4.54,11.98,19,1.4)
        elif hasvisual:
            ppt_actions(slide,actions[:3],.60,y,5.4,19 if page.get('id')=='gls-run' else 20,2.97)
            if page.get('code'): ppt_code(slide,page['code'],6.3,y,6.43,3.04)
            elif page.get('image'): ppt_picture(slide,page['image'],6.3,y,6.4,3.04)
            elif page.get('diagram'): ppt_diagram(slide,page['diagram'],6.12,y,6.45,2.98)
            else: ppt_table(slide,page['table'],6.1,y,6.64,3.05)
        else:
            ppt_actions(slide,actions[:4],.60,y,11.87,21,2.99)
        if page.get('checkpoint'): ppt_checkpoint(slide,page['checkpoint'])
        # Extra slides are purposeful demonstrations, never duplicate full paragraphs.
        detail_needed=(len(actions)> (2 if page.get('diagram') in ('route','checks','files') else 3 if hasvisual else 4)) or (bool(page.get('code')) and bool(page.get('table') or page.get('image'))) or bool(page.get('predict')) or (len(page.get('help',[]))>=3)
        if detail_needed:
            detail=new_slide(prs,page,i,True); notes(detail,page,i+1)
            if page.get('predict'):
                pred=page['predict']; ppt_rect(detail,.58,1.99,12.13,1.82,'FFF0E3','FFF0E3')
                ppt_text(detail,.82,2.20,11.6,.42,'PAUSE & PREDICT',16,ORANGE,True)
                ppt_text(detail,.82,2.70,11.6,1.08,pred['question'],26,NAVY,True)
                ppt_text(detail,.82,4.21,11.65,1.52,'Answer: '+pred['answer'],22,INK)
            elif page.get('code') and (page.get('table') or page.get('image')):
                if page.get('image'): ppt_picture(detail,page['image'],.65,2.05,8.05,3.9)
                else: ppt_table(detail,page['table'],.65,2.05,8.05,3.9)
                tail=actions[3:] or [t['word']+' = '+t['meaning'] for t in page.get('terms',[])]
                ppt_actions(detail,tail[:3],9.0,2.2,3.64,19,3.5)
            elif len(page.get('help',[]))>=3:
                short=[('Unknown module / missing file','Check the three names and both source lists. Guide page 7.'),
                       ('Reset / shift assertion fails','Check pin bits, reset polarity and expected values. Pages 5 and 9.'),
                       ('Documentation fails','Complete metadata, sections and pin keys. Keep YAML indentation. Page 8.'),
                       ('Layout / viewer / gate test fails','Open the first failing job. See GLS on pages 13–14 and help on page 15.')]
                for n,(problem,fix) in enumerate(short):
                    xx=.7+(n%2)*6.05; yy=2.06+(n//2)*1.84
                    ppt_rect(detail,xx,yy,5.72,1.58,WHITE)
                    ppt_text(detail,xx+.18,yy+.17,5.34,.42,problem,19,ORANGE,True)
                    ppt_text(detail,xx+.18,yy+.69,5.34,.73,fix,19,INK)
            elif page.get('table') and page.get('diagram') in ('files','checks'):
                ppt_table(detail,page['table'],.66,2.07,11.93,3.12)
                if page['diagram']=='checks': ppt_text(detail,.78,5.43,11.62,.55,actions[2],20,TEAL,True)
            else:
                start=2 if page.get('diagram') in ('route','checks','files') else 3 if hasvisual else 4
                ppt_actions(detail,actions[start:],.66,2.12,11.8,22,3.6)
            if page.get('checkpoint'): ppt_checkpoint(detail,page['checkpoint'])
    prs.core_properties.title=lesson['title']
    prs.core_properties.subject='Compact practical lesson, with editable flowcharts and guide references'
    prs.core_properties.author='EE6019 Research Project'
    prs.save(output_file('lesson.pptx'))
    return len(prs.slides)


def powershell_literal(path):
    return "'"+str(path).replace("'","''")+"'"


def export_pdf():
    REVIEW.mkdir(parents=True,exist_ok=True)
    docpath=powershell_literal(output_file('quickstart.docx')); pdfpath=powershell_literal(output_file('quickstart.pdf'))
    deckpath=powershell_literal(output_file('lesson.pptx')); deckpdf=powershell_literal(output_file('lesson.pdf'))
    renderpath=powershell_literal(REVIEW/'slides'); warningpath=powershell_literal(REVIEW/'slide-fit.json')
    script=f"""
$ErrorActionPreference = 'Stop'
$taskWordApp = New-Object -ComObject Word.Application
$taskWordApp.Visible = $false
$taskWordApp.DisplayAlerts = 0
$taskGuide = $null
try {{
    $taskGuide = $taskWordApp.Documents.Open({docpath}, $false, $false)
    $taskGuide.Repaginate()
    $null = $taskGuide.Fields.Update()
    $taskGuide.Save()
    $taskGuide.ExportAsFixedFormat({pdfpath}, 17)
    Write-Output ('Guide: ' + $taskGuide.ComputeStatistics(2) + ' pages')
}} finally {{
    if ($taskGuide) {{ $taskGuide.Close(0); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($taskGuide) }}
    $taskWordApp.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($taskWordApp)
}}
$taskOwnPpt = $false
try {{ $taskPptApp = [Runtime.InteropServices.Marshal]::GetActiveObject('PowerPoint.Application') }}
catch {{ $taskPptApp = New-Object -ComObject PowerPoint.Application; $taskOwnPpt = $true }}
$taskDeck = $null
try {{
    $taskDeck = $taskPptApp.Presentations.Open({deckpath}, $true, $false, $false)
    $taskDeck.SaveAs({deckpdf}, 32)
    New-Item -ItemType Directory -Force -Path {renderpath} | Out-Null
    $taskDeck.Export({renderpath}, 'PNG', 1600, 900)
    $taskWarnings = @()
    foreach ($taskSlide in $taskDeck.Slides) {{
        foreach ($taskShape in $taskSlide.Shapes) {{
            if ($taskShape.HasTextFrame -eq -1 -and $taskShape.TextFrame2.HasText -eq -1) {{
                $taskFrame = $taskShape.TextFrame2
                $taskRange = $taskFrame.TextRange
                $taskHeight = $taskShape.Height - $taskFrame.MarginTop - $taskFrame.MarginBottom
                $taskWidth = $taskShape.Width - $taskFrame.MarginLeft - $taskFrame.MarginRight
                if ($taskRange.BoundHeight -gt $taskHeight + 2 -or $taskRange.BoundWidth -gt $taskWidth + 2) {{
                    $taskWarnings += [pscustomobject]@{{slide=$taskSlide.SlideIndex;shape=$taskShape.Name;text=$taskRange.Text;height=$taskRange.BoundHeight;allowed_height=$taskHeight;width=$taskRange.BoundWidth;allowed_width=$taskWidth}}
                }}
            }}
        }}
    }}
    ConvertTo-Json -InputObject @($taskWarnings) -Depth 4 | Set-Content -LiteralPath {warningpath} -Encoding UTF8
    Write-Output ('Deck: ' + $taskDeck.Slides.Count + ' slides; text-fit warnings: ' + $taskWarnings.Count)
}} finally {{
    if ($taskDeck) {{ $taskDeck.Close(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($taskDeck) }}
    if ($taskOwnPpt -and $taskPptApp.Presentations.Count -eq 0) {{ $taskPptApp.Quit() }}
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($taskPptApp)
}}
"""
    path=REVIEW/'export.ps1'; path.write_text(script,encoding='utf-8-sig')
    result=subprocess.run(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(path)],capture_output=True,text=True,check=False)
    print(result.stdout)
    if result.returncode: raise RuntimeError(result.stderr)


def review_pdfs(lesson):
    import pymupdf as fitz
    report={}
    for fname,prefix in [('quickstart.pdf','guide'),('lesson.pdf','deck')]:
        path=output_file(fname)
        if not path.exists(): continue
        pdf=fitz.open(path); dest=REVIEW/prefix; dest.mkdir(parents=True,exist_ok=True)
        entries=[]
        for i,page in enumerate(pdf):
            pix=page.get_pixmap(matrix=fitz.Matrix(1.25,1.25),alpha=False)
            pix.save(str(dest/f'{i+1:02d}.png'))
            text=page.get_text()
            entries.append({'page':i+1,'words':len(text.split()),'first_lines':text.splitlines()[:5]})
        thumbs=[]
        for start in range(0,len(pdf),6):
            canvas=Image.new('RGB',(1320,1020),'#E7ECEE'); draw=ImageDraw.Draw(canvas)
            for j in range(start,min(start+6,len(pdf))):
                im=Image.open(dest/f'{j+1:02d}.png'); im.thumbnail((420,465))
                col=(j-start)%3; row=(j-start)//3; x=col*440+(440-im.width)//2; y=row*510+30
                canvas.paste(im,(x,y)); draw.text((col*440+15,row*510+8),f'{prefix} {j+1}',font=font(18,True),fill='#'+NAVY)
            out=REVIEW/f'{prefix}-contact-{start+1:02d}.png'; canvas.save(out); thumbs.append(str(out))
        report[prefix]={'pages':len(pdf),'entries':entries,'contacts':thumbs}
        if prefix=='guide' and len(pdf)!=len(lesson['pages']): report[prefix]['pagination_warning']='Expected one page per lesson section.'
    (REVIEW/'pdf-review.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({k:{'pages':v['pages'],'pagination_warning':v.get('pagination_warning')} for k,v in report.items()}))


def main():
    global OUTPUT_SUFFIX
    parser=argparse.ArgumentParser(); parser.add_argument('--export-pdf',action='store_true'); parser.add_argument('--qa',action='store_true'); parser.add_argument('--suffix',default='')
    args=parser.parse_args(); lesson=json.loads((BASE/'lesson.json').read_text(encoding='utf-8-sig'))
    OUTPUT_SUFFIX=args.suffix
    REVIEW.mkdir(parents=True,exist_ok=True); assets={}
    for kind in ('route','shift','mapping','files','checks'):
        path=REVIEW/f'diagram-{kind}.png'; draw_diagram(kind,path); assets[kind]=path
    write_doc(lesson,assets); count=write_slides(lesson)
    print(f'Built {len(lesson["pages"])} guide sections and {count} slides.')
    if args.export_pdf: export_pdf()
    if args.qa: review_pdfs(lesson)


if __name__=='__main__': main()
