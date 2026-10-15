"""Build the editable beginner DOCX from guide.md and genuine local figures."""
from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

HERE = Path(__file__).resolve().parent
NAVY = '19394A'
TEAL = '087E83'
BODY = '253E4B'
LIGHT = 'EAF4F3'

def shade(element, fill):
    props=element.get_or_add_tcPr() if element.tag==qn('w:tc') else element.get_or_add_pPr()
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);props.append(shd)

def field(paragraph, instruction):
    run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=instruction
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for el in [begin,code,sep,end]:run._r.append(el)

def link(paragraph,label,target):
    rel=paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True)
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);node.append(run)
    paragraph._p.append(node)

TOKEN=re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))')
def inline(paragraph,text):
    pos=0
    for m in TOKEN.finditer(text):
        paragraph.add_run(text[pos:m.start()]);token=m.group()
        if token.startswith('['):
            label,url=re.match(r'\[([^\]]+)\]\(([^)]+)\)',token).groups();link(paragraph,label,url)
        elif token.startswith('**'):paragraph.add_run(token[2:-2]).bold=True
        elif token.startswith('`'):
            r=paragraph.add_run(token[1:-1]);r.font.name='Consolas';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor.from_string(NAVY)
        else:paragraph.add_run(token[1:-1]).italic=True
        pos=m.end()
    paragraph.add_run(text[pos:])

def cells(line):
    # Vertical bars inside code are expressions, not table delimiters.
    line=line.strip().strip('|');parts=[];chunk='';code=False;escape=False
    for ch in line:
        if ch=='`' and not escape:code=not code
        if ch=='|' and not code and not escape:parts.append(chunk.strip());chunk=''
        else:chunk+=ch
        escape=ch=='\\' and not escape
    parts.append(chunk.strip());return parts

def add_table(doc,lines):
    rows=[cells(l) for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in r)]
    n=max(map(len,rows));table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    usable=6.87
    proportions={2:[.32,.68],3:[.26,.32,.42],4:[.18,.24,.24,.34],5:[.16,.2,.2,.22,.22]}.get(n,[1/n]*n)
    if n==3 and 'module' in rows[0][1].lower():
        proportions=[.22,.40,.38]
    keep_small=len(rows)<=8 and sum(len(c) for row in rows for c in row)<1400
    for col,fraction in zip(table.columns,proportions):col.width=Inches(usable*fraction)
    for i,row in enumerate(rows):
        cells_=table.add_row().cells
        trpr=table.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(cells_):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.width=Inches(usable*proportions[j])
            props=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for name,value in [('top','95'),('bottom','95'),('left','95'),('right','95')]:
                el=OxmlElement('w:'+name);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margin.append(el)
            props.append(margin)
            p=cell.paragraphs[0];p.style='Table text';inline(p,row[j] if j<len(row) else '')
            if keep_small:p.paragraph_format.keep_with_next=i<len(rows)-1
            if i==0:
                shade(cell._tc,TEAL)
                for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                for run in p._p.xpath('.//w:rPr'):
                    color=run.find(qn('w:color'))
                    if color is not None:color.set(qn('w:val'),'FFFFFF')
            elif i%2:shade(cell._tc,'F0F6F6')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def configure(doc):
    sec=doc.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.70);sec.right_margin=Inches(.70)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.30)
    normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor.from_string(BODY)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size,color in [('Title',34,NAVY),('Subtitle',18,TEAL),('Heading 1',25,NAVY),('Heading 2',17,TEAL),('Heading 3',13,NAVY)]:
        st=doc.styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(9);st.paragraph_format.keep_with_next=True
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    for name,size in [('Code',10.5),('Table text',10.5),('Figure caption',10),('Small print',10)]:
        if name not in doc.styles:doc.styles.add_style(name,1)
        st=doc.styles[name];st.base_style=normal;st.font.size=Pt(size)
        st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.05
    code=doc.styles['Code'];code.font.name='Consolas';code.font.color.rgb=RGBColor.from_string(NAVY)
    code.paragraph_format.space_before=Pt(0);code.paragraph_format.space_after=Pt(0);code.paragraph_format.line_spacing=1.0
    code.paragraph_format.keep_together=True
    cap=doc.styles['Figure caption'];cap.font.color.rgb=RGBColor.from_string('416575');cap.paragraph_format.keep_together=True
    header=sec.header.paragraphs[0];header.add_run('TINY TAPEOUT  |  A beginner’s laboratory manual').font.size=Pt(9)
    header.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('8 October 2026   •   Page ').font.size=Pt(9)
    field(foot,' PAGE ');foot.add_run(' of ');field(foot,' NUMPAGES ')
    settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.title='Your first Tiny Tapeout chip — a complete beginner guide'
    doc.core_properties.subject='Step-by-step counter project, free preparation and fabrication concepts'
    doc.core_properties.author='MD. KAMRUL ISLAM'
    doc.core_properties.keywords='Tiny Tapeout, beginner, Verilog, cocotb, ASIC, SKY130, tutorial'

def cover(doc):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(46)
    r=p.add_run('FROM YOUR FIRST BIT\nTO A CHECKED CHIP LAYOUT');r.bold=True;r.font.size=Pt(12);r.font.color.rgb=RGBColor.from_string(TEAL)
    doc.add_paragraph('Your first\nTiny Tapeout chip',style='Title')
    doc.add_paragraph('A complete beginner’s guide',style='Subtitle')
    doc.add_paragraph('Understand the circuit. Follow the steps. Read the evidence.',style='Normal')
    table=doc.add_table(rows=1,cols=1);cell=table.cell(0,0);shade(cell._tc,NAVY)
    p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(14);p.paragraph_format.space_after=Pt(14)
    r=p.add_run('One real example, explained all the way through.\nBits → Verilog → tests → GitHub Actions → layout → checks\nThen understand submission and fabrication.');r.font.size=Pt(13);r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph('Main route: a browser, a free GitHub account and a public repository.\nPractical endpoint: a checked layout and archived evidence, before any fabrication order.',style='Normal')
    doc.add_paragraph('Includes 30 genuine screenshot placements, worked waveforms, complete source files, troubleshooting and answered exercises.',style='Normal')
    p=doc.add_paragraph('EE6019 teaching material • 8 October 2026\nPrepared for student review • MD. KAMRUL ISLAM',style='Small print')
    p.paragraph_format.space_before=Pt(22)
    doc.add_page_break()
    p=doc.add_paragraph('Contents');p.style='Title'
    doc.add_paragraph('Use the linked contents or View → Navigation Pane in Word. The first time, read the parts in order.',style='Normal')
    p=doc.add_paragraph();field(p,' TOC \\o "1-2" \\h \\z \\u ')
    doc.add_paragraph('Worked explanations and answered exercises are in Part 1. The practical browser steps start in Part 2.',style='Small print')

def build():
    doc=Document();configure(doc);cover(doc)
    lines=(HERE/'guide.md').read_text(encoding='utf-8').splitlines();i=0;figure=0;code_count=0
    while i<len(lines):
        line=lines[i]
        if not line.strip() or line.strip()=='---':i+=1;continue
        fence=re.match(r'^(`{3,})(\w*)',line)
        if fence:
            marker=fence.group(1);i+=1;code_count+=1
            code_start=i
            code_end=next((j for j in range(i,len(lines)) if lines[j].strip()==marker),len(lines))
            while i<len(lines) and lines[i].strip()!=marker:
                p=doc.add_paragraph(style='Code');p.add_run(lines[i] or ' ');shade(p._p,'F0F4F5')
                # Keep useful code context after a heading and prevent a two-line tail.
                p.paragraph_format.keep_with_next=(i<code_start+4 or code_end-5<=i<code_end-1)
                i+=1
            if i==len(lines):raise ValueError('Unclosed code fence')
            # An empty paragraph after a full-page listing can create a blank page
            # before the next file's explicit page break. Space the final line instead.
            p.paragraph_format.space_after=Pt(8);i+=1;continue
        m=re.match(r'^(#{1,3}) (.+)',line)
        if m:
            heading=doc.add_heading(m.group(2),level=len(m.group(1)))
            if re.match(r'^A\.\d+ — ',m.group(2)) and not m.group(2).startswith('A.2 — '):
                heading.paragraph_format.page_break_before=True
            i+=1;continue
        m=re.match(r'^!\[(.*)\]\(([^)]+)\)$',line)
        if m:
            caption,path=m.groups();figure+=1
            with Image.open(HERE/path) as img:w,h=img.size
            width=min(6.87,6.87*4.95/(6.87*h/w)) if 6.87*h/w>4.95 else 6.87
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
            if path=='img/24-real-datasheet.png':
                # A whole datasheet needs a dedicated page at readable print size.
                width=min(6.87,8.8*w/h)
                p.paragraph_format.page_break_before=True
            shape=p.add_run().add_picture(str(HERE/path),width=Inches(width))
            shape._inline.docPr.set('descr',caption);shape._inline.docPr.set('title',f'Figure {figure}')
            p=doc.add_paragraph(style='Figure caption');inline(p,f'Figure {figure}. '+caption);i+=1;continue
        if line.lstrip().startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].lstrip().startswith('|'):rows.append(lines[i]);i+=1
            add_table(doc,rows);continue
        m=re.match(r'^(\s*)(?:([-*]) |(\d+)\. )(.+)',line)
        if m:
            indent,bullet,number,text=m.groups();p=doc.add_paragraph()
            p.paragraph_format.left_indent=Inches(.20+len(indent)/8);p.paragraph_format.first_line_indent=Inches(-.18)
            inline(p,('• ' if bullet else number+'. ')+text);i+=1;continue
        parts=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,3} |!\[|\||`{3,}|---$|\s*[-*] |\s*\d+\. )',lines[i]):parts.append(lines[i]);i+=1
        p=doc.add_paragraph();inline(p,' '.join(parts))
    target=HERE/'tiny-tapeout-beginner-guide.docx';doc.save(target)
    print(json.dumps({'docx':str(target),'figures':figure,'code_blocks':code_count,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},indent=2))

if __name__=='__main__':build()
