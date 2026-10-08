"""Build the complete editable Tiny Tapeout course using local, sourced content."""
from pathlib import Path
import json, math, hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
from beginner_visuals import draw_visual

HERE = Path(__file__).resolve().parent
GUIDE = HERE.parent / 'beginner-guide'
PAPER, INK, TEAL, CORAL = 'F7F4ED', '142C3A', '087F83', 'F06445'
WHITE, MIST, MUTED, RULE = 'FFFFFF', 'E3F0EE', '496471', 'C7D6D5'
FONT, DISPLAY, MONO = 'Aptos', 'Aptos Display', 'Consolas'
W, H = 13.333333, 7.5
SHA = '61ef06bed008e76000f9ae46d65182e476d15dd6'
FIT_PATH=HERE/'beginner-text-fit.json'
FIT=json.loads(FIT_PATH.read_text(encoding='utf-8-sig')) if FIT_PATH.exists() else {}
CURRENT_ID=''
LESSON_LINKS=[]

def colour(value): return RGBColor.from_string(value)

def box(slide,x,y,w,h,fill=WHITE,line=None,kind=MSO_SHAPE.RECTANGLE):
    obj=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    obj.fill.solid();obj.fill.fore_color.rgb=colour(fill)
    if line: obj.line.color.rgb=colour(line);obj.line.width=Pt(1)
    else: obj.line.fill.background()
    obj._element.spPr.append(OxmlElement('a:effectLst'))
    return obj

def text(slide,x,y,w,h,value,size=24,fill=INK,bold=False,font=FONT,align=PP_ALIGN.LEFT):
    name=f'Text[{x:.2f},{y:.2f}]: '+str(value).replace('\n',' ')[:65]
    size=FIT.get(CURRENT_ID+'|'+name,size)
    obj=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=obj.text_frame;tf.clear();tf.word_wrap=True
    tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=Inches(.02)
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.line_spacing=1.05;p.space_after=Pt(0)
        run=p.add_run();run.text=line;run.font.name=font;run.font.size=Pt(size)
        run.font.bold=bold;run.font.color.rgb=colour(fill)
    obj.name=name
    return obj

def line(slide,x1,y1,x2,y2,fill=TEAL,width=2):
    obj=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    obj.line.color.rgb=colour(fill);obj.line.width=Pt(width)
    return obj

def check(slide,value):
    if not value:return
    box(slide,.65,6.36,12.03,.67,MIST)
    text(slide,.82,6.38,11.68,.65,'CHECK  '+value,18,TEAL,True)

def numbered_body(slide,items,x,y,w,h,size=24):
    if not items:return
    gap=.18;row=(h-gap*(len(items)-1))/len(items)
    for i,value in enumerate(items):
        yy=y+i*(row+gap)
        box(slide,x,yy+.05,.31,.31,TEAL,kind=MSO_SHAPE.OVAL)
        text(slide,x+.015,yy+.095,.28,.22,str(i+1),11,WHITE,True,align=PP_ALIGN.CENTER)
        text(slide,x+.48,yy,w-.48,row,value,size)

def base(prs,data,number,total):
    global CURRENT_ID
    CURRENT_ID=data['id']
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=colour(PAPER)
    box(slide,0,0,.16,H,TEAL)
    text(slide,.65,.28,11.9,.29,'YOUR FIRST CHIP  /  '+data['section'].upper(),11,TEAL,True)
    text(slide,.65,.83,12.03,1.16,data['title'],36,INK,True,DISPLAY)
    box(slide,.65,7.09,12.03,.014,RULE)
    text(slide,.65,7.19,9.55,.19,'EE6019  •  8 October 2026  •  Sources and teaching notes in PowerPoint',9,MUTED)
    if data.get('sources'):
        ref=text(slide,9.93,7.17,1.47,.24,'Open source ↗',10,TEAL,True,align=PP_ALIGN.RIGHT)
        ref.text_frame.paragraphs[0].runs[0].hyperlink.address=data['sources'][0]
    text(slide,11.70,7.15,.97,.30,f'{number:03} / {total:03}',11,TEAL,True,align=PP_ALIGN.RIGHT)
    return slide

def explain(slide,data):
    body=data.get('body',[]);example=data.get('example')
    if example:
        numbered_body(slide,body,.65,2.12,7.10,3.90,24)
        box(slide,8.15,2.12,4.53,3.90,INK)
        text(slide,8.44,2.40,3.94,.31,'WORKED EXAMPLE',12,'71D5CC',True)
        text(slide,8.44,3.00,3.94,2.57,example,27,WHITE,bold=True)
    else:
        numbered_body(slide,body,.83,2.16,11.40,3.88,27)

def compare(slide,data):
    items=data.get('body',[])
    if len(items)>2:
        # Two rows keep four comparison statements comfortably readable.
        for i,value in enumerate(items):
            r,c=divmod(i,2);x=.65+c*6.17;y=2.12+r*2.04
            box(slide,x,y,5.86,1.82,WHITE)
            text(slide,x+.23,y+.22,5.40,1.34,value,25)
    else:
        for i,value in enumerate(items):
            x=.65+i*6.17;box(slide,x,2.12,5.86,3.90,INK if i==0 else MIST)
            text(slide,x+.28,2.46,5.30,3.18,value,29,WHITE if i==0 else INK)
    if data.get('example'):
        text(slide,.83,6.02,11.68,.33,data['example'],18,TEAL,True)

def code_slide(slide,data):
    code=data.get('code','');lines=code.splitlines()
    longest=max(map(len,lines),default=0)
    if longest<=56 and len(data.get('body',[]))<=3:
        box(slide,.65,2.12,7.70,3.76,INK)
        text(slide,.88,2.39,7.24,3.42,code,19,WHITE,font=MONO)
        numbered_body(slide,data.get('body',[]),8.56,2.18,4.11,3.61,21)
        if data.get('example'):
            text(slide,.88,5.95,11.62,.38,data['example'],18,TEAL,True)
    else:
        height=.40+len(lines)*.315
        box(slide,.65,2.08,12.03,height,INK)
        text(slide,.90,2.28,11.53,height-.31,code,18,WHITE,font=MONO)
        remaining=6.30-(2.08+height+.16)
        prose='  •  '.join(data.get('body',[])+(['EXAMPLE: '+data['example']] if data.get('example') else []))
        text(slide,.83,2.08+height+.16,11.77,remaining,prose,20)

def bits(slide,data):
    if data['id']=='C04':
        for x,value,label in [(1.00,'0','Logical low'),(7.05,'1','Logical high')]:
            box(slide,x,2.20,5.20,2.28,INK if value=='0' else TEAL)
            text(slide,x+.20,2.33,4.80,1.35,value,66,WHITE,True,MONO,PP_ALIGN.CENTER)
            text(slide,x+.20,3.78,4.80,.51,label,25,WHITE,True,align=PP_ALIGN.CENTER)
        text(slide,.83,4.73,11.90,.69,data['example'],28,TEAL,True)
        text(slide,.83,5.53,11.90,.80,'  •  '.join(data.get('body',[])),21)
        return
    value=int(data.get('value',5));weights=[128,64,32,16,8,4,2,1]
    for i,weight in enumerate(weights):
        x=.83+i*1.49;on=bool(value&weight)
        text(slide,x,2.20,1.25,.40,str(weight),23,TEAL,True,align=PP_ALIGN.CENTER)
        box(slide,x,2.71,1.25,1.22,CORAL if on else INK)
        text(slide,x,2.88,1.25,.90,'1' if on else '0',47,WHITE,True,MONO,PP_ALIGN.CENTER)
        text(slide,x,4.04,1.25,.40,f'bit {7-i}',18,MUTED,align=PP_ALIGN.CENTER)
    text(slide,.83,4.62,11.77,.85,data.get('example',f'{value} in decimal = {value:08b} in binary'),28,TEAL,True)
    text(slide,.83,5.53,11.90,.80,'  •  '.join(data.get('body',[])),21)

def flow(slide,data):
    stages=data.get('flow')
    if not stages:
        items=data.get('body',[])
        for i,value in enumerate(items):
            y=2.14+i*.94
            box(slide,.84,y+.02,.51,.51,TEAL,kind=MSO_SHAPE.OVAL)
            text(slide,.85,y+.075,.49,.36,str(i+1),20,WHITE,True,align=PP_ALIGN.CENTER)
            text(slide,1.67,y,10.98,.78,value,25)
            if i<len(items)-1:line(slide,1.10,y+.54,1.10,y+.94)
        if data.get('example'):text(slide,.83,5.94,11.77,.39,data['example'],20,TEAL,True)
        return
    n=len(stages)
    if n>4:
        for i,stage in enumerate(stages):
            row,col=divmod(i,3);x=.65+col*4.12;y=2.12+row*2.03
            card=box(slide,x,y,3.78,1.75,INK if i%2==0 else MIST)
            if stage.get('target_id'):LESSON_LINKS.append((card,stage['target_id']))
            fill=WHITE if i%2==0 else INK
            title=text(slide,x+.20,y+.19,3.38,.43,stage['name'],22,fill,True)
            detail=text(slide,x+.20,y+.80,3.38,.79,stage['detail'],20,fill)
            if stage.get('target_id'):
                LESSON_LINKS.extend([(title,stage['target_id']),(detail,stage['target_id'])])
            if stage.get('range'):
                span=text(slide,x+.20,y+.60,3.38,.25,stage['range'],11,fill,True)
                if stage.get('target_id'):LESSON_LINKS.append((span,stage['target_id']))
    else:
        width=(12.03-.28*(n-1))/n
        for i,stage in enumerate(stages):
            x=.65+i*(width+.28)
            box(slide,x,2.28,width,2.82,INK if i%2==0 else MIST)
            fill=WHITE if i%2==0 else INK
            text(slide,x+.20,2.56,width-.40,.77,stage['name'],24,fill,True)
            text(slide,x+.20,3.59,width-.40,1.25,stage['detail'],22,fill)
            if i<n-1:line(slide,x+width,3.75,x+width+.28,3.75)
    prose='  •  '.join(data.get('body',[])+([data['example']] if data.get('example') else []))
    if prose:text(slide,.83,5.16,11.77,1.11,prose,20,TEAL)

def trace(slide,data):
    rows=data.get('rows',[])
    if not rows:return explain(slide,data)
    headings=['Edge','rst_n','enable','Before','After','wrap']
    widths=[1.1,1.8,1.8,2.25,2.25,2.25];keys=['edge','rst_n','en','before','after','wrap']
    for r,values in enumerate([headings]+[[str(row[k]) for k in keys] for row in rows]):
        x=.65;y=2.06+r*.47
        for c,value in enumerate(values):
            fill=TEAL if r==0 else WHITE if r%2 else MIST
            box(slide,x,y,widths[c],.47,fill)
            text(slide,x+.08,y+.06,widths[c]-.16,.39,value,19,WHITE if r==0 else INK,r==0,align=PP_ALIGN.CENTER)
            x+=widths[c]
    text(slide,.83,5.70,11.67,.47,data.get('example','Predicted values at settled samples after each rising edge.'),22,TEAL,True)

def screenshot(slide,data):
    path=GUIDE/'img'/data['image']
    with Image.open(path) as img:w,h=img.size
    crop=data.get('crop',[0,0,0,0]);l,t,r,b=crop
    aspect=w*(1-l-r)/(h*(1-t-b))
    # Wide images get a large evidence panel; portrait pages need a focused crop.
    maxw,maxh=(12.03,3.00) if aspect>=2.4 else (7.64,3.70)
    pw=min(maxw,maxh*aspect);ph=pw/aspect
    x=.65+(maxw-pw)/2;y=2.13+(maxh-ph)/2
    pic=slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(pw),height=Inches(ph))
    pic.crop_left=l;pic.crop_top=t;pic.crop_right=r;pic.crop_bottom=b
    pic._element.nvPicPr.cNvPr.set('descr',data.get('caption',data['title']))
    if aspect>=2.4:
        text(slide,.83,5.22,11.80,.75,'  •  '.join(data.get('body',[])),21)
    else:
        numbered_body(slide,data.get('body',[]),8.50,2.21,4.30,3.63,21)
    # Attribution is present in the slide and full sources in notes.
    text(slide,.83,6.00,11.67,.33,data.get('caption',''),10,MUTED)

def exercise(slide,data):
    box(slide,.65,2.12,6.85,3.91,INK)
    text(slide,.92,2.38,6.31,.35,'TRY IT FIRST',12,'71D5CC',True)
    question='\n'.join(data.get('body',[]))+('\n\n'+data['question'] if data.get('question') else '')
    text(slide,.92,2.91,6.31,2.90,question,23,WHITE)
    box(slide,7.80,2.12,4.88,3.91,MIST)
    text(slide,8.07,2.38,4.34,.35,'WORKED ANSWER',12,TEAL,True)
    text(slide,8.07,3.02,4.34,2.59,data.get('answer',data.get('example','')),25,INK)

def cover(slide,data):
    # Native, editable chip illustration; conceptual rather than foundry evidence.
    box(slide,8.85,2.05,3.47,3.40,INK)
    for row in range(3):
        for col in range(3):
            box(slide,9.09+col*1.03,2.29+row*1.00,.87,.83,CORAL if(row,col)==(1,1)else MIST)
    for i in range(6):
        box(slide,8.57,2.49+i*.46,.28,.09,TEAL)
        box(slide,12.32,2.49+i*.46,.28,.09,TEAL)
    text(slide,.83,2.24,7.36,1.88,'Understand it.\nBuild it. Test it.',43,INK,True,DISPLAY)
    text(slide,.83,4.70,7.35,1.20,'One real 8-bit counter, from your first bit\nto a checked chip layout.',28)
    text(slide,8.84,5.75,3.50,.35,'Illustrative shared chip',14,MUTED)

LAYOUTS={'explain':explain,'compare':compare,'code':code_slide,'bits':bits,'flow':flow,'trace':trace,'screenshot':screenshot,'exercise':exercise,'cover':cover}

def build():
    LESSON_LINKS.clear()
    content=json.loads((HERE/'beginner-course.json').read_text(encoding='utf-8'))
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    prs.core_properties.title='Your first Tiny Tapeout chip — the complete beginner course'
    prs.core_properties.author='MD. KAMRUL ISLAM'
    prs.core_properties.subject='Step-by-step Verilog, tests, physical design and fabrication concepts'
    manifest=[]
    image_sources={Path(r['file']).name:r for r in json.loads((GUIDE/'screenshot-sources.json').read_text(encoding='utf-8'))}
    for number,data in enumerate(content,1):
        slide=base(prs,data,number,len(content))
        if not draw_visual(slide,data,box,text,line):
            LAYOUTS[data['layout']](slide,data)
        check(slide,data.get('check'))
        note=data.get('notes','')
        sources=data.get('sources',[])
        if data.get('image'):
            record=image_sources[data['image']]
            sources=list(dict.fromkeys(sources+[record.get('source_page',record.get('url','')),record.get('asset_url','')]))
            sources=[s for s in sources if s]
        slide.notes_slide.notes_text_frame.text=(
            'TEACHING NOTES\n'+note+'\n\nMANUAL REFERENCE\n'+data.get('guide_heading','')+
            '\n\nSOURCES / checked 8 October 2026\n'+'\n'.join(sources)+
            ('\n\nIMAGE PROVENANCE\n'+json.dumps(record,ensure_ascii=False,indent=2) if data.get('image') else ''))
        manifest.append({'slide':number,'id':data['id'],'title':data['title'],'section':data['section'],
                         'guide_heading':data.get('guide_heading'),'sources':sources,
                         'image':data.get('image'),'crop':data.get('crop'),'layout':data['layout']})
    by_id={spec['id']:slide for spec,slide in zip(content,prs.slides)}
    for shape,target in LESSON_LINKS:
        shape.click_action.target_slide=by_id[target]
    output=HERE/'tiny-tapeout-beginner-course.pptx';prs.save(output)
    (HERE/'beginner-slide-sources.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'slides':len(prs.slides),'screenshots':sum(bool(s.get('image'))for s in content),'output':str(output)},indent=2))

if __name__=='__main__':build()
